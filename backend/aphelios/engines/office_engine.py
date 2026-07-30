"""OfficeEngine – Word-/Excel-/PowerPoint-/PDF-Dokumente lesen und
analysieren (Alpha 1.6, erste Ausbaustufe).

Ausgelöst über ``/dokument <Pfad> [Frage]``: liest ein Office-Dokument oder
PDF von der Platte, extrahiert den Text (bei Excel: Zellen je
Tabellenblatt, bei PowerPoint: Text je Folie) und lässt Claude die Frage
dazu beantworten bzw. das Dokument zusammenfassen – ohne
``ANTHROPIC_API_KEY`` gibt es nur den rohen extrahierten Text (dasselbe
Fallback-Prinzip wie bei ``BrowserEngine``/``KnowledgeEngine``).

Unterstützte Formate (moderne XML-basierte Office-Open-XML-Formate + PDF –
bewusst NICHT die alten Binärformate .doc/.xls/.ppt, die python-docx/
openpyxl/python-pptx nicht lesen können):
    * ``.docx`` (Word) – python-docx
    * ``.xlsx`` (Excel) – openpyxl
    * ``.pptx`` (PowerPoint) – python-pptx
    * ``.pdf`` – pypdf

**Sicherheit:** Jede Anfrage läuft über das SecurityGate (CONFIRM) –
APHELIOS liest dabei den Inhalt einer beliebigen, vom Nutzer angegebenen
Datei UND schickt ihn (mit API-Key) an die Claude-API. Dieselbe Abwägung
wie bei ``VisionEngine`` (Screenshot) und ``BrowserEngine`` (Webseite): der
Inhalt kann beliebig sensibel sein.

**Bewusst NICHT in dieser ersten Ausbaustufe:**

- **Kein Erstellen/Schreiben neuer Dokumente.** Nur lesend – ein Dokument
  mit korrektem Layout/Formatierung neu zu erzeugen ist ein deutlich
  größerer, eigenständiger Scope (Vorlagen, Styling, Corporate Design …).
- **Keine Formel-Berechnung in Excel.** Zellen werden mit ihrem zuletzt
  gespeicherten Wert gelesen (``data_only=True``), Formeln selbst werden
  nicht neu ausgewertet.
- **Keine eingebetteten Bilder/Diagramme.** Nur Text (bei Excel: Zellwerte,
  bei PowerPoint: Text-Shapes) – für Bildinhalte gibt es bereits die
  ``VisionEngine`` (``/sieh`` auf einen Screenshot der geöffneten Datei).
- **Keine alten Binärformate** (.doc/.xls/.ppt vor Office 2007) – die
  genutzten Bibliotheken unterstützen nur die moderneren XML-Formate.

Bus-Schnittstelle:
    * ``office.request`` (in) – ``{id, text}`` (``text`` = Pfad, optional
      gefolgt von einer Frage)
    * ``chat.token`` / ``chat.response`` (out) – wie ``BrowserEngine``
"""

from __future__ import annotations

import asyncio
import re
from pathlib import Path

from aphelios.core.engine import BaseEngine
from aphelios.core.event_bus import Event
from aphelios.core.security import RiskLevel

OFFICE_PERSONA = """\
Du bist APHELIOS im Dokument-Analyse-Modus. Dir wird der extrahierte Text \
eines Office-Dokuments oder PDFs gezeigt. Beantworte die gestellte Frage \
ausschließlich anhand dieses Inhalts, bzw. fasse das Dokument knapp \
zusammen, wenn keine konkrete Frage gestellt wurde – ruhig, präzise, auf \
Deutsch. Ist die Antwort im gegebenen Text nicht enthalten, sag das ehrlich \
statt zu raten.
"""

#: Maximale Zeichen des extrahierten Dokumenttexts, die an Claude bzw. als
#: roher Fallback-Text gehen – siehe ``_MAX_PAGE_CHARS`` in browser_engine.py
#: für dieselbe Abwägung (lange Dokumente würden sonst unnötig viele Tokens
#: kosten bzw. den Chat fluten).
_MAX_DOCUMENT_CHARS = 8000

_SUPPORTED_EXTENSIONS = ("docx", "xlsx", "pptx", "pdf")

#: Sucht nach einer der unterstützten Datei-Endungen statt am ersten
#: Leerzeichen zu splitten – Windows-Pfade enthalten oft Leerzeichen (z. B.
#: "C:\\Users\\User\\Meine Dokumente\\Bericht.docx"), ein simples
#: Leerzeichen-Split würde den Pfad an der falschen Stelle abschneiden
#: (siehe ``_URL_LIKE_RE`` in browser_engine.py für das analoge Problem bei
#: URLs, dort ohne Leerzeichen-Risiko lösbar).
_PATH_RE = re.compile(
    r"^(.*?\.(?:" + "|".join(_SUPPORTED_EXTENSIONS) + r"))(?:\s+(.*))?$",
    re.IGNORECASE | re.DOTALL,
)


def _parse_document_request(text: str) -> tuple[str, str] | None:
    """Trennt eine ``/dokument``-Anfrage in Pfad und optionale Frage.

    Gibt ``None`` zurück, wenn keine unterstützte Datei-Endung erkennbar ist
    – der Aufrufer soll dann um einen konkreten Pfad bitten statt zu raten.
    """
    stripped = text.strip()
    if not stripped:
        return None
    match = _PATH_RE.match(stripped)
    if not match:
        return None
    path = match.group(1).strip()
    question = (match.group(2) or "").strip()
    return path, question


def _extract_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _extract_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), data_only=True, read_only=True)
    lines: list[str] = []
    for sheet in wb.worksheets:
        lines.append(f"## {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(cell) for cell in row if cell is not None]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def _extract_pptx(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    lines: list[str] = []
    for index, slide in enumerate(prs.slides, start=1):
        lines.append(f"## Folie {index}")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                lines.append(shape.text_frame.text.strip())
    return "\n".join(lines)


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _extract_document_text(path: Path) -> str:
    """Wählt die passende Extraktions-Funktion anhand der Datei-Endung.

    Läuft blockierend (läuft über ``asyncio.to_thread`` beim Aufrufer,
    genau wie mss/pytesseract in der ``VisionEngine``).
    """
    suffix = path.suffix.lower().lstrip(".")
    if suffix == "docx":
        return _extract_docx(path)
    if suffix == "xlsx":
        return _extract_xlsx(path)
    if suffix == "pptx":
        return _extract_pptx(path)
    if suffix == "pdf":
        return _extract_pdf(path)
    raise ValueError(f".{suffix} wird nicht unterstützt" if suffix else "Keine Datei-Endung erkannt")


class OfficeEngine(BaseEngine):
    """Liest Office-Dokumente/PDFs und lässt Claude ihren Inhalt beantworten."""

    name = "office"

    async def start(self) -> None:
        self._running = True
        self._client = None
        if self.config.has_anthropic:
            try:
                import anthropic

                self._client = anthropic.AsyncAnthropic(api_key=self.config.anthropic_api_key)
            except Exception:  # noqa: BLE001
                self._client = None
        self.bus.subscribe("office.request", self.handle)

    async def handle(self, event: Event) -> None:
        text = (event.data.get("text") or "").strip()
        request_id = event.data.get("id", "")

        parsed = _parse_document_request(text)
        if parsed is None:
            await self._reply(
                request_id,
                'Bitte einen Pfad zu einem Word-/Excel-/PowerPoint-Dokument oder PDF '
                'angeben, z. B. "/dokument C:\\Berichte\\Q3.docx Was ist das Fazit?"',
            )
            return
        path_str, question = parsed
        path = Path(path_str)

        if not await self._confirm_read(request_id, path_str):
            return

        if not path.exists():
            await self._reply(request_id, f"Datei nicht gefunden: {path_str}")
            return

        try:
            document_text = await asyncio.to_thread(_extract_document_text, path)
        except ImportError as exc:
            await self._reply(
                request_id,
                f'Dieses Dateiformat braucht ein zusätzliches Paket ({exc}) – '
                'pip install -e ".[office]", siehe docs/office.md.',
            )
            return
        except Exception as exc:  # noqa: BLE001
            self.log.exception("Dokument konnte nicht gelesen werden: %s", path_str)
            await self._reply(request_id, f"Dokument konnte nicht gelesen werden: {exc}"[:300])
            return

        document_text = document_text.strip()
        if not document_text:
            await self._reply(request_id, f'"{path.name}" geöffnet, aber kein Text gefunden.')
            return

        if self._client is None:
            await self._reply(
                request_id,
                f'Kein ANTHROPIC_API_KEY konfiguriert, deshalb nur roher Text aus "{path.name}":\n\n'
                f"{document_text[:_MAX_DOCUMENT_CHARS]}",
            )
            return

        await self._answer_with_claude(path.name, document_text, question, request_id)

    async def _confirm_read(self, request_id: str, path_str: str) -> bool:
        """Fragt vor JEDEM Datei-Zugriff das SecurityGate – der Inhalt kann
        beliebig sensibel sein und geht (mit API-Key) an die Claude-API."""
        allowed = await self.security.request(
            action="Dokument lesen",
            target=path_str,
            level=RiskLevel.CONFIRM,
            reason="APHELIOS liest den Inhalt dieser Datei und schickt ihn (mit API-Key) an Claude.",
        )
        if not allowed:
            await self._reply(request_id, "Abgelehnt.")
        return allowed

    async def _answer_with_claude(
        self, filename: str, document_text: str, question: str, request_id: str
    ) -> None:
        assert self._client is not None
        prompt = question or "Fasse den Inhalt dieses Dokuments kurz zusammen."
        user_message = (
            f"Dateiname: {filename}\n\n"
            f"Extrahierter Inhalt (ggf. gekürzt):\n{document_text[:_MAX_DOCUMENT_CHARS]}\n\n"
            f"Frage: {prompt}"
        )
        buffer: list[str] = []
        try:
            async with self._client.messages.stream(
                model=self.config.anthropic_model,
                max_tokens=800,
                system=OFFICE_PERSONA,
                messages=[{"role": "user", "content": user_message}],
            ) as stream:
                async for chunk in stream.text_stream:
                    buffer.append(chunk)
                    await self.emit("chat.token", {"id": request_id, "text": chunk})
        except Exception as exc:  # noqa: BLE001
            self.log.exception("Dokument-Analyse fehlgeschlagen")
            await self._reply(request_id, f"Analyse fehlgeschlagen: {str(exc)[:200]}")
            return
        await self.emit("chat.response", {"id": request_id, "text": "".join(buffer), "final": True})

    async def _reply(self, request_id: str, text: str) -> None:
        for word in text.split(" "):
            await self.emit("chat.token", {"id": request_id, "text": word + " "})
            await asyncio.sleep(0.02)
        await self.emit("chat.response", {"id": request_id, "text": text, "final": True})
